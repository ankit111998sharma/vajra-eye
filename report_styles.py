"""Shared Word formatting helpers matching Kurukshetra University MCA project rules."""
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor, Twips, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import OxmlElement
from pathlib import Path

FIG = Path(r"d:\MCA PROJECT KUK SEM 3\figures")
NAVY = RGBColor(0x1B, 0x36, 0x5D)
BLACK = RGBColor(0x00, 0x00, 0x00)


def set_run_font(run, name="Times New Roman", size=12, bold=False, italic=False, underline=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.underline = underline
    if color is not None:
        run.font.color.rgb = color


def _set_style_font(style, name, size, bold=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = BLACK
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), name)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    r = run._r
    r.append(fld1)
    r.append(instr)
    r.append(fld2)
    set_run_font(run, size=12)


def configure_section(section, different_first=False):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(p)
    # suppress header
    header = section.header
    header.is_linked_to_previous = False


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    _set_style_font(normal, "Times New Roman", 12)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    def ensure(name, base="Normal"):
        try:
            return styles.add_style(name, 1)
        except ValueError:
            return styles[name]

    ch = ensure("ChapterHeading")
    _set_style_font(ch, "Times New Roman", 20, bold=True)
    ch.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ch.paragraph_format.space_before = Pt(30)
    ch.paragraph_format.space_after = Pt(30)
    ch.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    ch.paragraph_format.page_break_before = True

    ph = ensure("ParaHeading")
    _set_style_font(ph, "Times New Roman", 14, bold=True)
    ph.font.underline = True
    ph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    ph.paragraph_format.space_before = Pt(12)
    ph.paragraph_format.space_after = Pt(12)
    ph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    sh = ensure("SubHeading")
    _set_style_font(sh, "Times New Roman", 13, bold=True)
    sh.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    sh.paragraph_format.space_before = Pt(10)
    sh.paragraph_format.space_after = Pt(8)
    sh.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    body = ensure("ReportBody")
    _set_style_font(body, "Times New Roman", 12)
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    body.paragraph_format.space_before = Pt(6)
    body.paragraph_format.space_after = Pt(6)
    body.paragraph_format.first_line_indent = Cm(0.75)

    cap = ensure("CaptionStyle")
    _set_style_font(cap, "Times New Roman", 11, bold=True)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(6)
    cap.paragraph_format.space_after = Pt(12)
    cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    code = ensure("CodeStyle")
    _set_style_font(code, "Courier New", 10)
    code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.space_before = Pt(0)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.first_line_indent = Cm(0)

    toc = ensure("TOCEntry")
    _set_style_font(toc, "Times New Roman", 12)
    toc.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    toc.paragraph_format.line_spacing = 1.15
    toc.paragraph_format.space_before = Pt(3)
    toc.paragraph_format.space_after = Pt(3)


def new_document():
    doc = Document()
    setup_styles(doc)
    configure_section(doc.sections[0])
    return doc


def chapter(doc, text):
    p = doc.add_paragraph(text, style="ChapterHeading")
    for r in p.runs:
        set_run_font(r, size=20, bold=True)
    return p


def heading(doc, text):
    p = doc.add_paragraph(text, style="ParaHeading")
    for r in p.runs:
        set_run_font(r, size=14, bold=True, underline=True)
    return p


def subhead(doc, text):
    p = doc.add_paragraph(text, style="SubHeading")
    for r in p.runs:
        set_run_font(r, size=13, bold=True)
    return p


def body(doc, text, indent=True):
    p = doc.add_paragraph(style="ReportBody")
    run = p.add_run(text)
    set_run_font(run, size=12)
    if not indent:
        p.paragraph_format.first_line_indent = Cm(0)
    return p


def bodies(doc, paragraphs):
    for t in paragraphs:
        if t:
            body(doc, t)


def caption(doc, text):
    p = doc.add_paragraph(text, style="CaptionStyle")
    for r in p.runs:
        set_run_font(r, size=11, bold=True)
    return p


def add_figure(doc, filename, cap, width_cm=14.5):
    path = FIG / filename
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.0
    if path.exists():
        run = p.add_run()
        run.add_picture(str(path), width=Cm(width_cm))
    else:
        run = p.add_run(f"[Figure missing: {filename}]")
        set_run_font(run, size=10, italic=True)
    caption(doc, cap)


def add_code_lines(doc, lines):
    for line in lines.split("\n"):
        p = doc.add_paragraph(style="CodeStyle")
        run = p.add_run(line if line != "" else " ")
        set_run_font(run, name="Courier New", size=10)


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:color"), "1B365D")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        set_run_font(run, size=11, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade_cell(cell, "1B365D")
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(str(val))
            set_run_font(run, size=10)
            set_cell_border(cell)
            if r_i % 2 == 1:
                shade_cell(cell, "EEF3F7")
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(1.25)
    if p.runs:
        p.runs[0].text = text
        set_run_font(p.runs[0], size=12)
    else:
        run = p.add_run(text)
        set_run_font(run, size=12)
    return p


def bullets(doc, items):
    for it in items:
        bullet(doc, it)


def center_line(doc, text, size=12, bold=False, italic=False, space_before=6, space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def formula(doc, text):
    """Centered displayed equation in Times New Roman (KUK body font)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_run_font(run, size=12, italic=True)
    return p


def page_break(doc):
    doc.add_page_break()
