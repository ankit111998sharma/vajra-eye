"""Front matter: cover, declaration, certificate, acknowledgement, synopsis, lists, TOC."""
from report_styles import (
    body, bodies, heading, subhead, bullets, center_line, page_break, add_table, caption
)
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm, RGBColor
from report_styles import set_run_font


def write_cover(doc):
    for _ in range(1):
        center_line(doc, "CENTRE FOR DISTANCE AND ONLINE EDUCATION", 14, True, space_before=0, space_after=0)
    center_line(doc, "KURUKSHETRA UNIVERSITY, KURUKSHETRA", 14, True, space_before=0, space_after=6)
    center_line(doc, "(Established by the State Legislature Act XII of 1956)", 11, False, True, space_before=0, space_after=18)
    center_line(doc, "PROJECT REPORT", 20, True, space_before=18, space_after=6)
    center_line(doc, "ON", 14, True, space_before=0, space_after=12)
    center_line(doc, "VAJRA-EYE", 22, True, space_before=6, space_after=6)
    center_line(doc, "An Edge-Based Intelligent Aerial Surveillance System", 16, True, space_before=0, space_after=0)
    center_line(doc, "for Weapon Detection", 16, True, space_before=0, space_after=12)
    center_line(doc, "Submitted in partial fulfilment of the requirements for the award of the degree of", 12, False, True, space_before=18, space_after=6)
    center_line(doc, "MASTER OF COMPUTER APPLICATIONS (MCA)", 14, True, space_before=6, space_after=4)
    center_line(doc, "Third Semester  |  Project Work — Assignment 1 Submission", 12, True, space_before=0, space_after=18)
    center_line(doc, "Submitted By", 12, True, space_before=12, space_after=4)
    center_line(doc, "ANKIT", 14, True, space_before=0, space_after=0)
    center_line(doc, "Student ID: 11023526", 12, False, space_before=0, space_after=0)
    center_line(doc, "Enrolment No.: 24DEOJULYMCA00190", 12, False, space_before=0, space_after=12)
    center_line(doc, "Under the Guidance of", 12, True, space_before=10, space_after=4)
    center_line(doc, "Dr. Pooja Sharma", 14, True, space_before=0, space_after=0)
    center_line(doc, "Professor, Indian Institute of Technology Mandi", 12, False, space_before=0, space_after=0)
    center_line(doc, "Kamand, Himachal Pradesh – 175005", 12, False, space_before=0, space_after=18)
    center_line(doc, "September 2026", 12, True, space_before=20, space_after=0)
    page_break(doc)


def write_declaration(doc):
    heading(doc, "DECLARATION BY THE STUDENT")
    bodies(doc, [
        "I, Ankit, Student ID 11023526, Enrolment No. 24DEOJULYMCA00190, pursuing the degree of Master of Computer Applications (MCA) from the Centre for Distance and Online Education, Kurukshetra University, Kurukshetra, hereby declare that the project entitled “VAJRA-EYE: An Edge-Based Intelligent Aerial Surveillance System for Weapon Detection” is my original work, carried out under the guidance of Dr. Pooja Sharma, Professor, Indian Institute of Technology Mandi (Himachal Pradesh).",
        "I further declare that this project has not been submitted earlier, either in part or in full, for the award of any degree or diploma to any university or institution. All sources of information used in this project have been duly acknowledged in the bibliography. I understand that plagiarism is not accepted under any circumstances as per the university project guidelines.",
        "I also declare that a soft copy of the project is to be submitted on the Learning Management System (LMS) together with this report, as required by the university.",
    ])
    body(doc, "Signature of the Student: ____________________", indent=False)
    body(doc, "Name: Ankit", indent=False)
    body(doc, "Course: MCA (Third Semester)", indent=False)
    body(doc, "University: Centre for Distance and Online Education, Kurukshetra University", indent=False)
    body(doc, "Date: 01 September 2026", indent=False)
    page_break(doc)


def write_certificate(doc):
    heading(doc, "CERTIFICATE FROM THE PROJECT GUIDE")
    body(doc, "(As per Annexure A of the University Guidelines)", indent=False)
    bodies(doc, [
        "This is to certify that this project entitled “VAJRA-EYE: An Edge-Based Intelligent Aerial Surveillance System for Weapon Detection” submitted in partial fulfilment of the degree of MASTER IN COMPUTER APPLICATION (MCA) to Kurukshetra University done by Mr. Ankit, Roll No. / Student ID 11023526, Enrolment No. 24DEOJULYMCA00190, is an authentic work carried out by him under my guidance. The matter embodied in this project work has not been submitted earlier for award of any degree or diploma to the best of my knowledge and belief.",
    ])
    body(doc, "", indent=False)
    body(doc, "Signature of the Student: ____________________          Signature of the Guide: ____________________", indent=False)
    body(doc, "Date: 01 September 2026", indent=False)
    body(doc, "", indent=False)
    body(doc, "Guide Name: Dr. Pooja Sharma", indent=False)
    body(doc, "Designation: Professor", indent=False)
    body(doc, "Full Address: Indian Institute of Technology Mandi, Kamand, Himachal Pradesh – 175005, India", indent=False)
    body(doc, "Qualification: Ph.D. (as per institutional records)", indent=False)
    body(doc, "Mobile: ____________________", indent=False)
    body(doc, "Email: ____________________", indent=False)
    page_break(doc)


def write_ack(doc):
    heading(doc, "ACKNOWLEDGEMENT")
    bodies(doc, [
        "In the acknowledgement page, as required by the university guidelines, the student recognises his indebtedness for guidance and assistance. I take this opportunity to express my sincere gratitude to all who have helped me during the project work.",
        "I am deeply grateful to my project guide, Dr. Pooja Sharma, Professor, Indian Institute of Technology Mandi, Himachal Pradesh, for her valuable guidance, critical comments, and encouragement throughout the study. Her insistence on a theoretically justified design and on honest reporting of limitations has shaped this report.",
        "I thank the Centre for Distance and Online Education, Kurukshetra University, Kurukshetra, and the faculty associated with the MCA programme, for providing the academic framework, the project guidelines, and the opportunity to undertake a full software-development project in the third semester.",
        "I acknowledge the authors of the research papers, standards, and open-source frameworks cited in the bibliography—particularly the communities behind YOLOv8/Ultralytics, Deep Java Library, OpenCV, Spring Boot, and ONNX Runtime—without whose publicly documented work this implementation would not have been feasible.",
        "I also thank my family and peers for their patience during the long hours of literature survey, coding, testing, and report writing. Any errors that remain are my own.",
        "Finally, I record that this work is an academic prototype for examination. It is not a certified defence product, and it must not be deployed against the public without lawful authority, operational testing, and institutional approval.",
    ])
    body(doc, "Ankit", indent=False)
    body(doc, "MCA, CDOE, Kurukshetra University", indent=False)
    page_break(doc)


def write_synopsis(doc):
    heading(doc, "SYNOPSIS / ABSTRACT OF THE PROJECT")
    body(doc, "(Submitted separately as required: approximately three to four pages covering the prescribed heads.)", indent=False)

    subhead(doc, "Name / Title of the Project")
    body(doc, "VAJRA-EYE: An Edge-Based Intelligent Aerial Surveillance System for Weapon Detection.")

    subhead(doc, "Statement about the Problem")
    bodies(doc, [
        "Conventional surveillance still depends on human operators watching live video. Fatigue, clutter, and the sheer number of streams cause missed events. Streaming high-resolution drone video to a distant server consumes bandwidth that border and rural links cannot guarantee, and it adds latency that a security incident cannot afford. There is therefore a need for a system that watches the video near the camera, understands whether a visible handheld weapon is present, and sends a small, protected alert to an officer instead of shipping the entire film.",
        "The problem is not only algorithmic. It is a software-engineering problem: the solution must run as a maintainable service, must be testable, must record evidence, and must fail safely. Many published demos solve only the detector and leave the rest to imagination. This project treats the whole pipeline as the problem.",
    ])

    subhead(doc, "Why is the particular topic chosen?")
    bodies(doc, [
        "The topic was chosen for four converging reasons. First, the MCA syllabus and the university guideline allow industry-relevant work in intelligent systems, networks, and software engineering; aerial analytics sits at that junction. Second, publicly discussed security needs of India—border watching, crowded-event safety, protection of plants and campuses—make weapon-aware cameras a legitimate academic case, provided ethics are respected. Third, most student projects in this area stop at a Python notebook; a Java edge service is a harder and more professional artefact. Fourth, the name Vajra-Eye is used as an academic identifier for an indigenous-style design: a sharp, local, digital sentry rather than a rented foreign cloud.",
        "The topic is also chosen because it forces the student to read real literature (YOLO, edge computing, cryptography, UAV imaging) rather than to clone a CRUD database. That reading is documented at length in Chapter 3.",
    ])

    subhead(doc, "Objective and Scope of the Project")
    bodies(doc, [
        "Objectives: (1) to design a layered architecture for aerial/elevated weapon alerting; (2) to implement real-time ingest of RTSP or file video in Java; (3) to gate deep inference by motion-based keyframing; (4) to run YOLOv8 through DJL/ONNX; (5) to emit encrypted alerts over SMS/email; (6) to document the full SDLC artefacts required by the university (DFD, ERD, PERT, tests, manual, data dictionary).",
        "Scope in: visible handheld weapon classes supported by the loaded model; edge node plus command APIs; laboratory evaluation; report and source listing. Scope out: concealed weapons, autonomous firing, swarm control, legal interception of private communications, and classified national systems. The project ends as a working prototype plus this report, not as a production accreditation.",
    ])

    subhead(doc, "Methodology (including a summary of the project)")
    bodies(doc, [
        "An iterative waterfall-plus-prototype SDLC is used: literature and requirements, design diagrams, implementation of Spring services, integration of OpenCV and DJL, testing, and report. The runtime summary is: capture a frame; convert to grey and blur; difference against the previous frame; if motion exceeds a threshold, run YOLOv8; if a weapon-like class exceeds confidence 0.85, encrypt a JSON alert and notify. Only keyframes and events leave the node in the intended operational mode.",
        "Hardware and software to be used are listed next. Testing uses functional test cases, a laboratory confusion matrix, latency observation, and resource comparison with all-frame inference. Contribution is a complete, documented, JVM-native edge pipeline occupying a gap identified in the literature survey.",
    ])

    subhead(doc, "Hardware and Software to be used")
    add_table(doc,
              ["Item", "Specification / Product"],
              [
                  ["Processor", "Intel i5/i7 or edge-class CPU (Jetson optional)"],
                  ["RAM", "16 GB recommended (8 GB minimum)"],
                  ["Storage", "SSD, 20 GB free plus model files"],
                  ["Camera", "USB camera, IP RTSP camera, or recorded drone video"],
                  ["OS", "Windows 10/11 64-bit or Linux 64-bit"],
                  ["Language", "Java 17"],
                  ["Frameworks", "Spring Boot, DJL, ONNX Runtime, OpenCV"],
                  ["Build", "Apache Maven"],
                  ["Model", "YOLOv8n ONNX"],
                  ["Notify", "SMTP and/or Twilio SMS API"],
                  ["Report tools", "Microsoft Word (this document), Python for figures"],
              ])

    subhead(doc, "Testing Technologies used")
    body(doc, "JUnit-style unit tests for pure functions (encryption helpers, threshold logic), integration runs against recorded videos, manual exploratory tests of RTSP disconnects, and quantitative tallies of true/false detections on a labelled laboratory set. No single commercial test-automation suite is mandated; the methodology is recorded in Chapter 12 with test cases and a sample confusion matrix.")

    subhead(doc, "What contribution would the project make?")
    bodies(doc, [
        "Academically, the project contributes a full MCA-grade system document that other students can audit. Technically, it contributes a reusable pattern: motion-gated DJL inference inside Spring. Socially, if later institutionalised with lawful data and human command, it could shorten the time between a visible weapon appearing in a camera and an officer knowing about it, without hauling raw video across a weak network.",
        "The contribution is not a claim of battlefield readiness. It is a claim of a working, explainable prototype aligned with the university’s demand for software development, soft copy, and a properly formatted report.",
    ])
    page_break(doc)


def write_lists(doc):
    heading(doc, "LIST OF ABBREVIATIONS")
    add_table(doc,
              ["Abbreviation", "Expansion"],
              [
                  ["AES", "Advanced Encryption Standard"],
                  ["API", "Application Programming Interface"],
                  ["CDOE", "Centre for Distance and Online Education"],
                  ["CIA", "Confidentiality, Integrity, Availability"],
                  ["CNN", "Convolutional Neural Network"],
                  ["CPU", "Central Processing Unit"],
                  ["CCTV", "Closed-Circuit Television"],
                  ["DFD", "Data Flow Diagram"],
                  ["DJL", "Deep Java Library"],
                  ["DRDO", "Defence Research and Development Organisation"],
                  ["ERD", "Entity Relationship Diagram"],
                  ["FN/FP/TN/TP", "False/True Negative/Positive"],
                  ["FPS", "Frames Per Second"],
                  ["GCM", "Galois/Counter Mode"],
                  ["HOG", "Histogram of Oriented Gradients"],
                  ["HTTP/S", "Hypertext Transfer Protocol (Secure)"],
                  ["IVA", "Intelligent Video Analytics"],
                  ["JVM", "Java Virtual Machine"],
                  ["KUK", "Kurukshetra University, Kurukshetra"],
                  ["LMS", "Learning Management System"],
                  ["MCA", "Master of Computer Applications"],
                  ["NMS", "Non-Maximum Suppression"],
                  ["ONNX", "Open Neural Network Exchange"],
                  ["PERT", "Programme Evaluation and Review Technique"],
                  ["RBAC", "Role-Based Access Control"],
                  ["REST", "Representational State Transfer"],
                  ["RGB", "Red Green Blue"],
                  ["RTSP", "Real-Time Streaming Protocol"],
                  ["SDLC", "Software Development Life Cycle"],
                  ["SMS", "Short Message Service"],
                  ["SMTP", "Simple Mail Transfer Protocol"],
                  ["SSD", "Single Shot MultiBox Detector (or Solid State Drive, by context)"],
                  ["SVM", "Support Vector Machine"],
                  ["TLS", "Transport Layer Security"],
                  ["UAV", "Unmanned Aerial Vehicle"],
                  ["UML", "Unified Modeling Language"],
                  ["YOLO", "You Only Look Once"],
              ])
    caption(doc, "Table A: List of abbreviations used in this report.")

    heading(doc, "LIST OF FIGURES")
    figs = [
        "2.1 Logical layer stack of Vajra-Eye",
        "4.1 Use-case diagram (conceptual)",
        "4.2 Context-level DFD (Level 0)",
        "4.3 Level-1 data flow diagram",
        "4.4 Entity relationship diagram",
        "4.5 Core class collaboration",
        "5.1 Simplified PERT network",
        "5.2 Project Gantt chart (20 weeks)",
        "5.3 Iterative SDLC adopted",
        "6.1 Layered architecture",
        "6.2 Deployment view",
        "6.3 Sequence: detection and alert",
        "6.4 Camera node state model",
        "8.1 Detection pipeline flow",
        "8.2 Adaptive keyframe selection (Mt > θ)",
        "9.1 Defence-in-depth security layers",
        "10.1 Operator login (schematic)",
        "10.2 Live surveillance dashboard (schematic)",
        "10.3 Threat alert console (schematic)",
        "10.4 Keyframe evidence review (schematic)",
        "10.5 Admin policy panel (schematic)",
        "10.6 Audit log viewer (schematic)",
        "12.1 Confusion matrix (laboratory)",
        "12.2 Performance impact of motion keyframing",
    ]
    for i, f in enumerate(figs, 1):
        body(doc, f"Figure {f}", indent=False)

    heading(doc, "LIST OF TABLES")
    tabs = [
        "3.1 Qualitative comparison of related system patterns",
        "3.2 Research gaps versus project response",
        "4.1 Functional requirements (selected)",
        "4.2 Non-functional requirements",
        "5.1 Work-breakdown with PERT times (weeks)",
        "6.1 Technology stack",
        "7.1 Hardware used in the laboratory",
        "9.1 Security controls mapped to threats",
        "11.1 Cost elements (illustrative INR)",
        "12.1 Sample functional test cases",
        "12.2 Laboratory confusion counts",
        "A.1 Data dictionary",
        "A.2 Guide details",
    ]
    for t in tabs:
        body(doc, f"Table {t}", indent=False)
    page_break(doc)


def write_toc(doc):
    heading(doc, "TABLE OF CONTENTS")
    items = [
        "Cover Page",
        "Declaration by the Student",
        "Certificate from the Project Guide (Annexure A)",
        "Acknowledgement",
        "Synopsis / Abstract (3–4 pages as prescribed)",
        "List of Abbreviations",
        "List of Figures",
        "List of Tables",
        "Table of Contents",
        "CHAPTER 1  Introduction",
        "    1.1 Background of the study",
        "    1.2 Evolution of surveillance practice",
        "    1.3 Need for intelligent aerial surveillance in India",
        "    1.4 Project overview",
        "    1.5 Problem statement and definition of the problem",
        "    1.6 Why this topic was chosen",
        "    1.7 Objectives of the project",
        "    1.8 Scope and limitations of scope",
        "    1.9 Expected contribution",
        "    1.10 Organisation of the report",
        "CHAPTER 2  Theoretical Background and Conceptual Framework",
        "    2.1–2.9 Image, motion, CNN detectors, edge, crypto, architecture theory",
        "    2.10–2.18 Formal framework: paradigm, keyframe maths, YOLOv8, DJL, AES-GCM, metrics",
        "CHAPTER 3  Related Work and Literature Survey",
        "    3.1 Survey method",
        "    3.2 Traditional CCTV",
        "    3.3 Intelligent video frameworks",
        "    3.4 Evolution of object detectors",
        "    3.5 Weapon detection literature",
        "    3.6 UAV surveillance",
        "    3.7 Edge intelligence",
        "    3.8 Motion and keyframes",
        "    3.9 Java, DJL, enterprise hosting",
        "    3.10 Security and privacy",
        "    3.11 Indian indigenous context",
        "    3.12 Comparative analysis",
        "    3.13 Thematic synthesis",
        "    3.14 Research gaps",
        "    3.15 Positioning of Vajra-Eye",
        "    3.16 Additional reviewed works",
        "    3.17 How literature shaped decisions",
        "    3.18 Chapter summary",
        "    3.19–3.50 Extended commentaries, Indian context, open problems",
        "    3.51–3.82 YOLO lineage, datasets, weather, tracking, law, annotated notes",
        "CHAPTER 4  System Analysis and Design vis-à-vis User Requirements",
        "CHAPTER 5  System Planning (PERT / Gantt / SDLC)",
        "CHAPTER 6  System Architecture and Process Logic of Modules",
        "CHAPTER 7  Methodology, Implementation, Hardware and Software",
        "CHAPTER 8  Motion-Based Keyframing and Detection Pipeline",
        "CHAPTER 8A Fault-Tolerant Edge Implementation (memory, threads, RTSP, JVM)",
        "CHAPTER 9  Alerting, Security Mechanism and Controls",
        "CHAPTER 10 Input and Output Screen Design",
        "CHAPTER 11 Cost and Benefit Analysis",
        "CHAPTER 12 Testing Methodology, Test Report and Performance",
        "CHAPTER 13 System Maintenance and Evaluation",
        "CHAPTER 14 Applications and National Significance",
        "CHAPTER 15 User / Operational Manual",
        "CHAPTER 16 Conclusion, Limitations and Future Scope",
        "ANNEXURE I    Background of the organisation (academic setting)",
        "ANNEXURE II   Data dictionary",
        "ANNEXURE III  Guide details",
        "ANNEXURE IV   Printout of the code sheet (selected source)",
        "ANNEXURE V    References / Bibliography / Websites",
    ]
    for it in items:
        body(doc, it, indent=False)
    page_break(doc)
